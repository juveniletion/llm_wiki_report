# -*- coding: utf-8 -*-
"""
template_docx.py — **解析 Word 格式的报告模板**（赛题 5.1.1）

赛题原文：
    系统需支持解析 **Word 格式**的报告模板，识别固定章节标题和动态数据占位符
    （如 {{产品名称}}、{{本月材料成本}}、{{环比变动率}} 等）

它和「导出 Word」是两件事
------------------------
    **解析 Word 模板**（本模块）  输入侧：给一个 `.docx`，读懂章节与占位符，填数
    **导出 Word**（export_doc）   输出侧：把已生成的报告转成 Word/PDF

此前只有输出侧，`report_build.py` 填数时读的是 `.md` 模板，
那份 `.docx` **从未被程序打开过**——本条要求因此不达标。

两种产物，各有取舍
------------------
| | 产出 | 样式 |
|:---|:---|:---|
| 走 `.md` 模板（默认） | 先出 md 报告，再 `export_doc` 转 Word/PDF | **我们自己的排版**（封面/目录/图表） |
| 走 `.docx` 模板（本模块） | 直接填这个 docx → 出 Word | **保留模板原有样式**（用户的字体/表格样式） |

⇒ 用户给什么模板，产出就长什么样。这是"解析模板"该有的语义。

⚠️ 三个实测确认过的坑
--------------------
① **占位符可能被 Word 拆进多个 run**。
   实测赛题那份模板**没有**拆（108 处全完整落在单 run 里），
   但**用户自己做的模板经常拆**——Word 会因拼写检查、格式变化把
   `{{产品名称}}` 切成 `{{产品` + `名称}}`。
   所以本模块先试"逐 run 替换"（保留全部样式），
   失败再退回"合并段落替换"。

② **`}}%` 要紧跟处理**（模板里 31 处形如 `{{产量环比}}%`）。
   值统一带 `%`，若占位符后面已经有个 `%`，就剥掉值里那个，
   否则出现 `+5.45%%`。这与 `.md` 侧的 `fill()` 是同一规则。

③ **空值不能留 `{{}}`**。缺失的值填 `—`（与 md 侧一致），
   或者按调用方要求**整段删除**（可选占位符如 `{{重点分析段落}}`）。
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))
from console_io import ensure_utf8_stdout  # noqa: E402

ensure_utf8_stdout()

HERE = Path(__file__).resolve().parent
WIKI_ROOT = HERE.parent

_PH = re.compile(r"\{\{([^{}]+)\}\}")
_OL = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl"


# ---------------------------------------------------------------------------
# 读：识别章节标题与占位符
# ---------------------------------------------------------------------------

@dataclass
class TemplateInfo:
    path: Path
    sections: List[Tuple[int, str]] = field(default_factory=list)   # (大纲级, 标题)
    placeholders: Set[str] = field(default_factory=set)
    n_tables: int = 0
    n_paragraphs: int = 0

    def describe(self) -> str:
        lines = [f"模板：{self.path.name}",
                 f"  段落 {self.n_paragraphs} · 表格 {self.n_tables} · "
                 f"占位符 {len(self.placeholders)} 种"]
        if self.sections:
            lines.append("  识别到的章节（按大纲级别）：")
            for lv, title in self.sections[:24]:
                lines.append("    " + "  " * min(lv, 4) + f"[{lv}] {title}")
            if len(self.sections) > 24:
                lines.append(f"    … 另有 {len(self.sections) - 24} 个标题")
        return "\n".join(lines)


def _iter_paragraphs(doc: Any):
    """遍历文档的**全部**段落：正文 + 表格单元格（含嵌套表格）。"""
    for p in doc.paragraphs:
        yield p

    def walk_table(t: Any):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
                for nt in cell.tables:      # 嵌套表格
                    yield from walk_table(nt)

    for t in doc.tables:
        yield from walk_table(t)


def read_template(path: Path) -> TemplateInfo:
    """
    解析一个 `.docx` 模板：**章节标题 + 占位符**。

    章节的判据是**大纲级别**（`w:outlineLvl`），不是"看起来像标题"——
    实测赛题那份模板的标题段落 `style=None`，但带 `outlineLvl=1/2/3`。
    按大纲级别认才准，也才与 Word 的目录一致。
    """
    import docx
    doc = docx.Document(str(path))
    info = TemplateInfo(path=path, n_tables=len(doc.tables))

    for p in _iter_paragraphs(doc):
        info.n_paragraphs += 1
        txt = p.text.strip()
        if not txt:
            continue
        for m in _PH.finditer(p.text):
            info.placeholders.add(m.group(1).strip())
        # 大纲级别
        pPr = p._element.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr")
        lv = None
        if pPr is not None:
            e = pPr.find(_OL)
            if e is not None:
                lv = int(e.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"))
        # ⚠️ 标题里**可以带占位符**——实测「四、重点产品专项分析 — {{产品名称}}」
        #    就是这样。早先加了 `not _PH.search(txt)` 想排除"纯占位符段落"，
        #    结果把这一整章漏掉了（21 个标题报成 20 个）。
        #    改为：只要带大纲级别就算标题，占位符照留（它会被填成产品名）。
        if lv is not None:
            info.sections.append((lv, txt))

    info.sections.sort(key=lambda x: x[0])
    return info


# ---------------------------------------------------------------------------
# 填：把值写进模板的占位符
# ---------------------------------------------------------------------------

def _sub_text(text: str, values: Dict[str, Any],
              missing: List[str]) -> str:
    """
    在**一段文本**上替换占位符。规则与 `.md` 侧的 `fill()` 一致：
      · 缺失 → 填 `—`（显式，不留空）
      · `}}%` 紧跟百分号时剥掉值末尾的 `%`（否则 `+5.45%%`）
    """
    def rep(m: re.Match) -> str:
        name = m.group(1).strip()
        if name not in values:
            missing.append(name)
            return "—"
        v = "" if values[name] is None else str(values[name])
        # 占位符后面紧跟 `%` 且值也以 `%` 结尾 → 剥掉值里的那个
        if text[m.end():m.end() + 1] == "%" and v.endswith("%"):
            v = v[:-1]
        return v

    return _PH.sub(rep, text)


def fill_docx(template: Path, values: Dict[str, Any], out_path: Path,
              blank_keys: Optional[Set[str]] = None) -> Tuple[Path, List[str]]:
    """
    把 `values` 填进 docx 模板，**保留模板原有样式**。

    :param blank_keys: 这些占位符**整段清空**而不是填 `—`。
        用于"可选段落"（如 `{{重点分析段落}}`）——本就没有告警时，
        留一个 `—` 很怪，整段空着才对。
    :return: `(输出路径, 未填的占位符名列表)`

    ⚠️ 替换策略分两层（见模块 docstring 的坑 ①）：
       ① 先**逐 run 替换** —— 占位符完整落在单个 run 里时走这条，
          `run` 是 Word 的样式单位，不动它就不会丢格式。
       ② 不成再**合并整段替换** —— 占位符被拆进多个 run 时，
          把段落文本拼起来替换，结果写回第一个 run。
          代价：该段落内部的局部样式（如半句加粗）会丢。
          **这是"跨 run 拆分"的必然代价**，不是实现偷懒。
    """
    import docx
    blank_keys = blank_keys or set()
    doc = docx.Document(str(template))
    missing: List[str] = []

    def do_paragraph(par: Any) -> None:
        runs = par.runs
        if not runs:
            return
        # 整段文本里没有 `{{` 就跳过（绝大多数段落）
        full = "".join(r.text for r in runs)
        if "{{" not in full:
            return

        # 可选段落：整段就是这一个占位符且要求清空 → 置空整段
        only = _PH.fullmatch(full.strip())
        if only and only.group(1).strip() in blank_keys:
            for r in runs:
                r.text = ""
            return

        # ① 逐 run
        changed = False
        for r in runs:
            if "{{" in r.text:
                new = _sub_text(r.text, values, missing)
                if new != r.text:
                    r.text = new
                    changed = True
        if changed:
            return

        # ② 合并整段（跨 run 拆分的情形）
        new_full = _sub_text(full, values, missing)
        if new_full != full:
            runs[0].text = new_full
            for r in runs[1:]:
                r.text = ""

    for p in _iter_paragraphs(doc):
        do_paragraph(p)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))

    # 去重但保持顺序：同一个占位符在多处缺失时只报一次
    seen: List[str] = []
    for m in missing:
        if m not in seen:
            seen.append(m)
    return out_path, seen


# ---------------------------------------------------------------------------
# 自测 / 检视
# ---------------------------------------------------------------------------

def _main() -> int:
    ap = argparse.ArgumentParser(description="解析 Word 报告模板（赛题 5.1.1）")
    ap.add_argument("--inspect", metavar="DOCX", help="解析并打印章节与占位符")
    ap.add_argument("--template", metavar="DOCX", help="用它填充并输出")
    ap.add_argument("--out", help="输出路径")
    a = ap.parse_args()

    if a.inspect:
        info = read_template(Path(a.inspect))
        print(info.describe())
        print("\n  占位符（前 20 个）：")
        for i, k in enumerate(sorted(info.placeholders)[:20], 1):
            print(f"    {i:>2}. {{{{{k}}}}}")
        print(f"    … 共 {len(info.placeholders)} 种")
        return 0

    if a.template:
        tpl = Path(a.template)
        if not tpl.exists():
            print(f"模板不存在：{tpl}")
            return 1
        info = read_template(tpl)
        print(f"解析到 {len(info.placeholders)} 种占位符、{len(info.sections)} 个章节标题")
        # 演示：把所有占位符填成「样例值」
        values = {k: f"「{k}」" for k in info.placeholders}
        out = Path(a.out or (tpl.parent / (tpl.stem + "_filled.docx")))
        p, miss = fill_docx(tpl, values, out)
        print(f"✅ 已填充：{p}（未填 {len(miss)} 个）")
        return 0

    ap.print_help()
    return 1


if __name__ == "__main__":
    sys.exit(_main())
